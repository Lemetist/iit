from pathlib import Path
from html import escape
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import subprocess,re
p=Path(__file__).resolve().parent
out=(p/'run.txt').read_text()
counts=[int(x) for x in re.findall(r'\((\d+)\)',out)]
assert sum(counts)==750 and counts==[90,78,76,75,71,65,81,68,75,71]
for seed in [1,2,42,100]:
    run=subprocess.check_output([str(p/'task1'),str(seed)],text=True)
    assert sum(map(int,re.findall(r'\((\d+)\)',run)))==750
assert subprocess.check_output([str(p/'task1'),'12'],text=True)==out
(p/'output.html').write_text('<meta charset="utf-8"><style>body{background:#fff;color:#111;padding:25px}pre{font:16px/1.5 monospace}</style><pre>'+escape('$ ./task1 12\n'+out)+'</pre>')
# Блок-схема с циклами и обработкой правой границы.
im=Image.new('RGB',(1200,1500),'white'); dr=ImageDraw.Draw(im)
f=ImageFont.truetype('/usr/share/fonts/google-noto-vf/NotoSans[wght].ttf',24)
def box(y,text,shape='rect'):
    rect=(250,y,950,y+85)
    if shape=='oval': dr.ellipse(rect,outline='black',width=3)
    elif shape=='diamond': dr.polygon([(600,y-10),(970,y+42),(600,y+95),(230,y+42)],outline='black',width=3)
    else: dr.rectangle(rect,outline='black',width=3)
    dr.multiline_text((600,y+42),text,font=f,fill='black',anchor='mm',align='center')
def arrow(points,label=None):
    dr.line(points,fill='black',width=3)
    x,y=points[-1]; px,py=points[-2]
    if y>py: tip=[(x,y),(x-7,y-13),(x+7,y-13)]
    elif x>px: tip=[(x,y),(x-13,y-7),(x-13,y+7)]
    else: tip=[(x,y),(x+13,y-7),(x+13,y+7)]
    dr.polygon(tip,fill='black')
    if label: dr.text(points[0],label,font=f,fill='black')
items=[(20,'Начало','oval'),(140,'N = 750, K = 10, a = 100, b = 110\nsrand(seed), i = 0','rect'),(280,'i < N?','diamond'),(420,'numbers[i] = a + (b − a) · rand()/RAND_MAX\ni = i + 1','rect'),(590,'counts = 0, i = 0','rect'),(730,'i < N?','diamond'),(880,'j = int((numbers[i] − a) / ((b − a)/K))\nЕсли j >= K, то j = K − 1','rect'),(1020,'counts[j] = counts[j] + 1\ni = i + 1','rect'),(1190,'Вывод гистограммы, суммы частот\nи максимального отклонения','rect'),(1350,'Конец','oval')]
for y,t,s in items: box(y,t,s)
for y1,y2 in [(105,140),(225,270),(375,420),(675,720),(825,880),(965,1020),(1275,1350)]: arrow([(600,y1),(600,y2)])
dr.text((620,385),'Да',font=f,fill='black'); dr.text((620,835),'Да',font=f,fill='black')
arrow([(950,460),(1070,460),(1070,245),(600,245),(600,270)])
arrow([(230,322),(100,322),(100,560),(600,560),(600,590)],'Нет')
arrow([(950,1060),(1090,1060),(1090,700),(600,700),(600,720)])
arrow([(230,772),(100,772),(100,1155),(600,1155),(600,1190)],'Нет')
im.save(p/'flow.png')

d=Document(); s=d.sections[0]; s.top_margin=s.bottom_margin=Cm(1.7); s.left_margin=Cm(2.3); s.right_margin=Cm(1.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(12 if name=='Normal' else 16)
    st.paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.12

def para(t): d.add_paragraph(t)
def h(t): d.add_heading(t,1)
def code(t,size=9):
    q=d.add_paragraph(); q.paragraph_format.line_spacing=1; q.paragraph_format.space_after=Pt(0)
    r=q.add_run(t); r.font.name='Liberation Mono'; r.font.size=Pt(size)
def page(): d.add_page_break()
para('Образовательная организация __________________________')
para('Кафедра ____________________________________________')
for _ in range(3): para('')
d.add_paragraph('Отчёт по заданию 1', 'Title')
d.add_paragraph('Разработка алгоритмов генерации псевдослучайных последовательностей','Title')
para('Задача 1 • Вариант 12')
for _ in range(3): para('')
para('Выполнил(а) ____________________    Группа __________')
para('Проверил(а) ________________________________________')
para('Город ____________________    2026')
page(); h('Постановка задачи')
para('Сформировать массив из 750 псевдослучайных вещественных чисел в интервале [100; 110]. Разбить диапазон на 10 равных интервалов, подсчитать число попаданий и построить горизонтальную гистограмму. Оценить равномерность распределения по полученным частотам.')
h('Алгоритм решения')
for t in ['1. Инициализировать rand() функцией srand(). При обычном запуске seed берётся из времени; аргумент командной строки задаёт фиксированное значение для проверки.', '2. Заполнить массив по формуле x = 100 + 10 · (double(rand()) / RAND_MAX). Вещественное деление сохраняет дробную часть.', '3. Обнулить 10 счётчиков. Для каждого элемента вычислить номер интервала j = int((x − 100) / 1). Значение 110 отнести к последнему интервалу, заменив j = 10 на j = 9.', '4. Вывести интервалы [100; 101), …, [109; 110] и по одной звёздочке на каждое попавшее число. Все интервалы, кроме последнего, открыты справа.', '5. Сравнить частоты с ожидаемым числом 750 / 10 = 75; вычислить наибольшее абсолютное отклонение и его долю от 75. Проверить сумму частот.']: para(t)
h('Контрольный тест')
para('Команда ./task1 12. Получены частоты: 90, 78, 76, 75, 71, 65, 81, 68, 75, 71. Сумма равна 750. Максимальное отклонение от 75 равно 15 (20%). Повторный запуск с seed = 12 дал тот же вывод. При seed = 1, 2, 42 и 100 сумма частот также равна 750.')
para('Проверка границ по формуле: 100 → интервал 0; 101 → интервал 1; 109 → интервал 9; 110 → интервал 9 после коррекции. Последовательность rand() может различаться между реализациями стандартной библиотеки.')
h('Сборка и запуск'); code('g++ -std=c++17 -Wall -Wextra -pedantic task1.cpp -o task1\n./task1\n./task1 12',10)
page(); h('Скриншот результата работы')
d.add_picture(str(p/'screenshot.png'),width=Cm(17))
para('Рисунок 1 — вывод фактически выполненной программы ./task1 12, отображённый в браузере. Одна звёздочка соответствует одному числу.')
h('Выводы по работе')
para('Создан массив из 750 псевдослучайных чисел и построена горизонтальная гистограмма. В контрольной выборке частоты находятся в пределах 65–90 при ожидаемом значении 75. Гистограмма не показывает выраженной концентрации в одной части диапазона; различия частот естественны для конечной выборки. Эти наблюдения сами по себе не доказывают равномерность распределения и независимость чисел.')
para('Фиксированное начальное значение обеспечивает воспроизводимость в одной реализации rand(). Инициализация временем позволяет менять последовательность между запусками, начатыми в разные секунды.')
page(); h('Блок схема программы'); d.add_picture(str(p/'flow.png'),width=Cm(16.4))
page(); h('Программа на псевдокоде'); code('НАЧАЛО\n  N ← 750; K ← 10; a ← 100; b ← 110\n  h ← (b − a) / K\n  seed ← аргумент запуска или текущее время\n  Инициализировать генератор значением seed\n  ДЛЯ i ОТ 0 ДО N − 1\n    numbers[i] ← a + (b − a) · rand() / RAND_MAX\n    // Деление вещественное\n  КОНЕЦ ЦИКЛА\n  counts[0..K−1] ← 0\n  ДЛЯ КАЖДОГО x ИЗ numbers\n    j ← целая часть ((x − a) / h)\n    ЕСЛИ j ≥ K, ТО j ← K − 1\n    counts[j] ← counts[j] + 1\n  КОНЕЦ ЦИКЛА\n  expected ← N / K\n  total ← 0; maxDeviation ← 0\n  ДЛЯ i ОТ 0 ДО K − 1\n    Вывести границы интервала i\n    Вывести counts[i] звёздочек и значение counts[i]\n    total ← total + counts[i]\n    maxDeviation ← max(maxDeviation, |counts[i] − expected|)\n  КОНЕЦ ЦИКЛА\n  Вывести total, expected, maxDeviation\n  Вывести 100 · maxDeviation / expected процентов\nКОНЕЦ',10)
page(); h('Листинг программы на C++'); code((p/'task1.cpp').read_text(),9)
d.save(p/'Отчёт_задача_1_вариант_12.docx')
print('Report saved; repeatability and count checks passed')
